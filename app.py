import json
import os
import re
import uuid
from datetime import date, datetime, timedelta
from io import BytesIO

import bleach
import boto3
import openpyxl
import xlrd
from botocore.exceptions import BotoCoreError, ClientError
from botocore.config import Config
from docx import Document
from dotenv import load_dotenv
from flask import Flask, flash, jsonify, redirect, render_template, request, send_file, url_for

import ai
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import func, inspect, or_

load_dotenv()

app = Flask(__name__)
app.config["SECRET_KEY"] = os.getenv("SECRET_KEY", "change-this-secret")
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL", "sqlite:///worklog.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

db = SQLAlchemy(app)

ALLOWED_TAGS = set(bleach.sanitizer.ALLOWED_TAGS) | {
    "p", "br", "h1", "h2", "h3", "h4", "blockquote", "pre", "code",
    "ul", "ol", "li", "strong", "em", "u", "s", "a"
}
ALLOWED_ATTRIBUTES = {"a": ["href", "title", "target", "rel"]}


class Task(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    description = db.Column(db.Text, default="")
    status = db.Column(db.String(20), default="todo", nullable=False)
    priority = db.Column(db.String(20), default="medium", nullable=False)
    due_date = db.Column(db.Date, nullable=True)
    parent_id = db.Column(db.Integer, db.ForeignKey("task.id"), nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow, nullable=False)
    notes = db.relationship("Note", backref="task", cascade="all, delete-orphan", order_by="Note.updated_at.desc()")
    attachments = db.relationship("Attachment", backref="task", cascade="all, delete-orphan", order_by="Attachment.created_at.desc()")
    children = db.relationship("Task", backref=db.backref("parent", remote_side=[id]), order_by="Task.created_at", cascade="all, delete-orphan")


class Note(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    content = db.Column(db.Text, default="")
    task_id = db.Column(db.Integer, db.ForeignKey("task.id"), nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow, nullable=False)
    attachments = db.relationship("Attachment", backref="note", cascade="all, delete-orphan", order_by="Attachment.created_at.desc()")


class Folder(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    parent_id = db.Column(db.Integer, db.ForeignKey("folder.id"), nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    parent = db.relationship("Folder", remote_side=[id], backref=db.backref("children", cascade="all, delete-orphan"))
    files = db.relationship("Attachment", backref="folder", order_by="Attachment.created_at.desc()")


class Attachment(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    original_name = db.Column(db.String(255), nullable=False)
    object_key = db.Column(db.String(500), nullable=False, unique=True)
    content_type = db.Column(db.String(120), default="application/octet-stream")
    size = db.Column(db.Integer, default=0)
    task_id = db.Column(db.Integer, db.ForeignKey("task.id"), nullable=True)
    note_id = db.Column(db.Integer, db.ForeignKey("note.id"), nullable=True)
    folder_id = db.Column(db.Integer, db.ForeignKey("folder.id"), nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)


class Schedule(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    start_at = db.Column(db.DateTime, nullable=False)
    end_at = db.Column(db.DateTime, nullable=True)
    description = db.Column(db.Text, default="")
    task_id = db.Column(db.Integer, db.ForeignKey("task.id"), nullable=True)
    completed = db.Column(db.Boolean, default=False, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    task = db.relationship("Task", backref=db.backref("schedules", order_by="Schedule.start_at"))


class KnowledgeDoc(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    filename = db.Column(db.String(255), default="")
    content = db.Column(db.Text, default="")
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    chunks = db.relationship("KnowledgeChunk", backref="doc", cascade="all, delete-orphan", order_by="KnowledgeChunk.chunk_index")


class KnowledgeChunk(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    doc_id = db.Column(db.Integer, db.ForeignKey("knowledge_doc.id"), nullable=False)
    chunk_index = db.Column(db.Integer, nullable=False)
    text = db.Column(db.Text, nullable=False)
    embedding = db.Column(db.Text, nullable=False)  # JSON 数组


class ObjectStorage:
    def __init__(self):
        self.bucket = os.getenv("OSS_BUCKET")
        self.endpoint = os.getenv("OSS_ENDPOINT")
        self.client = None
        if self.bucket and self.endpoint and os.getenv("OSS_ACCESS_KEY_ID") and os.getenv("OSS_SECRET_ACCESS_KEY"):
            self.client = boto3.client(
                "s3",
                endpoint_url=self.endpoint,
                aws_access_key_id=os.getenv("OSS_ACCESS_KEY_ID"),
                aws_secret_access_key=os.getenv("OSS_SECRET_ACCESS_KEY"),
                region_name=os.getenv("OSS_REGION", "oss-cn-hangzhou"),
                config=Config(
                    signature_version="s3v4",
                    s3={"addressing_style": "virtual", "payload_signing_enabled": False},
                    request_checksum_calculation="when_required",
                    response_checksum_validation="when_required",
                ),
            )

    @property
    def enabled(self):
        return self.client is not None

    def upload(self, file_storage, key):
        if not self.enabled:
            raise RuntimeError("OSS 尚未配置")
        data = file_storage.read()
        self.client.put_object(Bucket=self.bucket, Key=key, Body=data, ContentLength=len(data), ContentType=file_storage.content_type or "application/octet-stream")
        return len(data)

    def download(self, key):
        response = self.client.get_object(Bucket=self.bucket, Key=key)
        return response["Body"].read()

    def delete(self, key):
        if self.enabled:
            self.client.delete_object(Bucket=self.bucket, Key=key)

    def upload_bytes(self, data, key, content_type):
        if self.enabled:
            self.client.put_object(Bucket=self.bucket, Key=key, Body=data, ContentLength=len(data), ContentType=content_type)

    def download_optional(self, key):
        if not self.enabled:
            return None
        try:
            return self.download(key)
        except ClientError as error:
            if error.response.get("Error", {}).get("Code") in {"NoSuchKey", "NoSuchBucket", "404"}:
                return None
            raise


storage = ObjectStorage()


def database_path():
    return os.path.join(app.instance_path, "worklog.db")


def restore_database():
    if not storage.enabled:
        return
    data = storage.download_optional(os.getenv("OSS_DATABASE_KEY", "worklog/worklog.db"))
    if data:
        os.makedirs(app.instance_path, exist_ok=True)
        with open(database_path(), "wb") as database_file:
            database_file.write(data)


def sync_database():
    if storage.enabled and os.path.exists(database_path()):
        with open(database_path(), "rb") as database_file:
            storage.upload_bytes(database_file.read(), os.getenv("OSS_DATABASE_KEY", "worklog/worklog.db"), "application/x-sqlite3")


# ---- 知识库 RAG ----
CHUNK_SIZE = 500
CHUNK_OVERLAP = 60


def chunk_text(text):
    text = re.sub(r"\s+\n", "\n", text.strip())
    chunks = []
    start = 0
    while start < len(text):
        piece = text[start:start + CHUNK_SIZE]
        if len(piece) < CHUNK_SIZE // 3 and chunks:
            chunks[-1] += piece
        else:
            chunks.append(piece)
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return [c for c in chunks if c.strip()]


def embed_document(doc):
    """为文档分块并写入向量（无向量接口时降级为空向量，检索走关键词匹配）。"""
    KnowledgeChunk.query.filter_by(doc_id=doc.id).delete()
    chunks = chunk_text(doc.content)
    vectors = []
    try:
        for start in range(0, len(chunks), 16):
            vectors.extend(ai.embed(chunks[start:start + 16]))
    except ai.EmbedUnavailable:
        vectors = []
    for index, piece in enumerate(chunks):
        vector = vectors[index] if index < len(vectors) else []
        doc.chunks.append(KnowledgeChunk(chunk_index=index, text=piece, embedding=json.dumps(vector)))
    db.session.commit()


def search_knowledge(question, top_k=5):
    """检索与问题最相关的知识块：优先向量相似，无向量时降级关键词匹配。"""
    query_tokens = ai.tokenize(question)
    try:
        query_vector = ai.embed([question])[0]
    except ai.EmbedUnavailable:
        query_vector = None
    results = []
    for chunk in KnowledgeChunk.query.all():
        vector = json.loads(chunk.embedding) if chunk.embedding else []
        if query_vector and vector:
            score = sum(a * b for a, b in zip(query_vector, vector))
        else:
            score = ai.keyword_score(query_tokens, ai.tokenize(chunk.text))
        results.append((chunk, chunk.doc, score))
    results.sort(key=lambda item: item[2], reverse=True)
    return results[:top_k]


def ask_knowledge(question):
    """RAG：检索知识库 → 组装上下文 → AI 生成回答。"""
    if not KnowledgeDoc.query.count():
        return {"answer": "知识库还是空的，请先上传制度文件或管理办法。", "sources": []}
    hits = search_knowledge(question)
    context_blocks = []
    for chunk, doc, _ in hits:
        context_blocks.append(f"【{doc.title}】\n{chunk.text}")
    prompt = (
        "你是企业制度助理。请优先根据下面的知识库资料回答问题；"
        "如果资料中没有相关内容，请如实说明知识库中没有找到，并基于常识简要回答。回答使用简体中文。\n\n"
        "=== 知识库资料 ===\n" + "\n\n".join(context_blocks) + "\n\n=== 问题 ===\n" + question
    )
    answer = ai.chat([{"role": "user", "content": prompt}], temperature=0.2)
    sources = [{"doc": doc.title, "snippet": chunk.text[:120]} for chunk, doc, _ in hits]
    return {"answer": answer, "sources": sources}


def sanitize_html(value):
    return bleach.clean(value or "", tags=ALLOWED_TAGS, attributes=ALLOWED_ATTRIBUTES, protocols=["http", "https", "mailto"], strip=True)


def parse_date(value):
    if not value:
        return None
    return datetime.strptime(value, "%Y-%m-%d").date()


def parse_datetime(value):
    return datetime.strptime(value, "%Y-%m-%dT%H:%M") if value else None


def plain_text(value):
    return re.sub(r"<[^>]+>", " ", value or "")


def root_task_for(task):
    current = task
    while current and current.parent:
        current = current.parent
    return current


def extract_attachment_text(data, filename, content_type):
    extension = os.path.splitext(filename.lower())[1]
    if extension == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        document = Document(BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if extension == ".xlsx" or content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
        return "\n".join(
            "\t".join(str(cell) for cell in row if cell is not None)
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows(values_only=True)
        )
    if extension == ".xls" or content_type == "application/vnd.ms-excel":
        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        return "\n".join(
            "\t".join(str(cell) for cell in worksheet.row_values(row_index) if cell != "")
            for worksheet in workbook.sheets()
            for row_index in range(worksheet.nrows)
        )
    if (content_type or "").startswith("text/"):
        return data.decode("utf-8", errors="ignore")
    return ""


def migrate_schema():
    inspector = inspect(db.engine)
    columns = {column["name"] for column in inspector.get_columns("attachment")}
    if "folder_id" not in columns:
        db.session.execute(db.text("ALTER TABLE attachment ADD COLUMN folder_id INTEGER REFERENCES folder(id)"))
        db.session.commit()
    schedule_columns = {column["name"] for column in inspector.get_columns("schedule")}
    if "completed" not in schedule_columns:
        db.session.execute(db.text("ALTER TABLE schedule ADD COLUMN completed BOOLEAN NOT NULL DEFAULT 0"))
        db.session.commit()
    task_columns = {column["name"] for column in inspector.get_columns("task")}
    if "parent_id" not in task_columns:
        db.session.execute(db.text("ALTER TABLE task ADD COLUMN parent_id INTEGER REFERENCES task(id)"))
        db.session.commit()


@app.context_processor
def inject_counts():
    return {"pending_count": Task.query.filter(Task.status != "done").count() if db.engine else 0, "oss_enabled": storage.enabled, "now": datetime.utcnow(), "timedelta": timedelta}


@app.route("/")
def index():
    status = request.args.get("status", "")
    query = request.args.get("q", "").strip()
    task_query = Task.query
    schedule_query = Schedule.query
    if status in {"todo", "done"}:
        task_query = task_query.filter_by(status=status)
        schedule_query = schedule_query.filter(Schedule.completed.is_(status == "done"))
    if query:
        task_query = task_query.filter(or_(Task.title.ilike(f"%{query}%"), Task.description.ilike(f"%{query}%")))
        schedule_query = schedule_query.filter(or_(Schedule.title.ilike(f"%{query}%"), Schedule.description.ilike(f"%{query}%")))
    tasks = task_query.order_by(Task.created_at.desc()).limit(20).all()
    schedules = schedule_query.order_by(Schedule.start_at).limit(20).all()

    # 合并任务与日程为统一清单
    items = []
    for task in tasks:
        root_task = root_task_for(task)
        items.append({
            "kind": "task",
            "id": task.id,
            "title": task.title,
            "description": task.description or "",
            "done": task.status == "done",
            "priority": task.priority,
            "parent_title": root_task.title if root_task and root_task.id != task.id else "",
            "time_label": task.due_date.strftime("%m月%d日截止") if task.due_date else "",
            "sort_time": datetime.combine(task.due_date, datetime.min.time()) if task.due_date else task.created_at,
        })
    for item in schedules:
        label = item.start_at.strftime("%m月%d日 %H:%M")
        if item.end_at:
            label += " - " + item.end_at.strftime("%H:%M")
        items.append({
            "kind": "schedule",
            "id": item.id,
            "title": item.title,
            "description": item.description or "",
            "done": bool(item.completed),
            "priority": None,
            "time_label": label,
            "sort_time": item.start_at,
        })
    pending = sorted((i for i in items if not i["done"]), key=lambda i: i["sort_time"])
    finished = sorted((i for i in items if i["done"]), key=lambda i: i["sort_time"], reverse=True)
    items = pending + finished
    shown_pending, shown_finished = len(pending), len(finished)

    notes = Note.query.order_by(Note.updated_at.desc()).limit(8).all()
    total_pending = Task.query.filter(Task.status != "done").count() + Schedule.query.filter(Schedule.completed.is_(False)).count()
    total_done = Task.query.filter(Task.status == "done").count() + Schedule.query.filter(Schedule.completed.is_(True)).count()
    today = date.today()
    today_tasks = Task.query.filter(Task.due_date == today, Task.status != "done").order_by(Task.priority).all()
    day_start = datetime.combine(today, datetime.min.time())
    day_end = datetime.combine(today, datetime.max.time())
    today_schedules = Schedule.query.filter(Schedule.start_at.between(day_start, day_end)).order_by(Schedule.start_at).all()
    today_items = []
    for task in today_tasks:
        root_task = root_task_for(task)
        today_items.append({"kind": "task", "id": task.id, "title": task.title, "done": False,
                            "time_label": "截止今天", "priority": task.priority,
                            "parent_title": root_task.title if root_task and root_task.id != task.id else ""})
    for item in today_schedules:
        label = item.start_at.strftime("%H:%M")
        if item.end_at:
            label += " - " + item.end_at.strftime("%H:%M")
        today_items.append({"kind": "schedule", "id": item.id, "title": item.title, "done": bool(item.completed),
                            "time_label": label, "priority": None})
    return render_template("index.html", items=items, notes=notes, current_status=status, query=query,
                           total_pending=total_pending, total_done=total_done,
                           shown_pending=shown_pending, shown_finished=shown_finished,
                           today_items=today_items, today_label=today.strftime("%m月%d日"))


@app.get("/api/reminders")
def api_reminders():
    """到点日程提醒：返回最近 5 分钟内开始且未完成的日程，前端轮询弹窗提醒。"""
    now = datetime.now()
    window_start = now - timedelta(minutes=5)
    items = (
        Schedule.query
        .filter(Schedule.completed.is_(False), Schedule.start_at <= now, Schedule.start_at >= window_start)
        .order_by(Schedule.start_at)
        .all()
    )
    return jsonify([{
        "id": item.id,
        "title": item.title,
        "time_label": item.start_at.strftime("%H:%M") + (f" - {item.end_at.strftime('%H:%M')}" if item.end_at else ""),
        "task": item.task.title if item.task else "",
        "description": (item.description or "")[:100],
    } for item in items])


@app.route("/schedule", methods=["GET", "POST"])
def schedule():
    if request.method == "POST":
        start_at = parse_datetime(request.form.get("start_at"))
        end_at = parse_datetime(request.form.get("end_at"))
        if not start_at:
            flash("请填写开始时间", "error")
        else:
            item = Schedule(title=request.form["title"].strip(), start_at=start_at, end_at=end_at, description=request.form.get("description", "").strip(), task_id=request.form.get("task_id", type=int) or None)
            db.session.add(item)
            db.session.commit()
            sync_database()
            flash("日程已创建", "success")
        return redirect(url_for("schedule", week=request.form.get("week")))
    selected = request.args.get("week")
    try:
        week_start = datetime.strptime(selected, "%Y-%m-%d") if selected else datetime.now()
    except ValueError:
        week_start = datetime.now()
    week_start = (week_start - timedelta(days=week_start.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    week_end = week_start + timedelta(days=7)
    items = Schedule.query.filter(Schedule.start_at >= week_start, Schedule.start_at < week_end).order_by(Schedule.start_at).all()
    days = [week_start + timedelta(days=index) for index in range(7)]
    time_slots = [f"{hour:02d}:{minute:02d}" for hour in range(6, 25) for minute in (0, 30)]
    return render_template("schedule.html", items=items, days=days, time_slots=time_slots, week_start=week_start, previous_week=week_start - timedelta(days=7), next_week=week_start + timedelta(days=7), tasks=Task.query.order_by(Task.title).all())


@app.post("/schedule/<int:schedule_id>/delete")
def delete_schedule(schedule_id):
    item = db.get_or_404(Schedule, schedule_id)
    db.session.delete(item)
    db.session.commit()
    sync_database()
    flash("日程已删除", "success")
    return redirect(request.referrer or url_for("schedule"))


@app.post("/schedule/<int:schedule_id>/status")
def update_schedule_status(schedule_id):
    item = db.get_or_404(Schedule, schedule_id)
    item.completed = request.form.get("completed") == "1"
    db.session.commit()
    sync_database()
    return redirect(request.referrer or url_for("schedule"))


@app.route("/search")
def global_search():
    query = request.args.get("q", "").strip()
    task_results = []
    note_results = []
    schedule_results = []
    file_results = []
    if query:
        pattern = f"%{query}%"
        task_results = Task.query.filter(or_(Task.title.ilike(pattern), Task.description.ilike(pattern))).all()
        note_results = Note.query.filter(or_(Note.title.ilike(pattern), Note.content.ilike(pattern))).all()
        schedule_results = Schedule.query.filter(or_(Schedule.title.ilike(pattern), Schedule.description.ilike(pattern))).order_by(Schedule.start_at).all()
        candidates = Attachment.query.filter(Attachment.original_name.ilike(pattern)).all()
        for attachment in Attachment.query.order_by(Attachment.created_at.desc()).all():
            if attachment in candidates or not storage.enabled:
                continue
            try:
                content = extract_attachment_text(storage.download(attachment.object_key), attachment.original_name, attachment.content_type)
                if query.lower() in content.lower():
                    candidates.append(attachment)
            except (BotoCoreError, ClientError, UnicodeError, ValueError, KeyError, OSError):
                continue
        file_results = candidates
    return render_template("search.html", query=query, task_results=task_results, note_results=note_results, schedule_results=schedule_results, file_results=file_results)


@app.get("/attachments/<int:attachment_id>/view")
def view_attachment(attachment_id):
    attachment = db.get_or_404(Attachment, attachment_id)
    query = request.args.get("q", "").strip()
    try:
        content = extract_attachment_text(storage.download(attachment.object_key), attachment.original_name, attachment.content_type)
    except (BotoCoreError, ClientError, UnicodeError, ValueError, KeyError, OSError) as error:
        flash(f"附件读取失败：{error}", "error")
        return redirect(request.referrer or url_for("global_search", q=query))
    return render_template("attachment_detail.html", attachment=attachment, content=content, query=query)


@app.route("/tasks/new", methods=["GET", "POST"])
def new_task():
    """任务创建统一在导图页完成，此入口重定向过去。"""
    return redirect(url_for("mindmap"))


@app.route("/tasks/<int:task_id>")
def task_detail(task_id):
    task = db.get_or_404(Task, task_id)
    return render_template("task_detail.html", task=task, available_files=Attachment.query.order_by(Attachment.original_name).all())


@app.route("/tasks/<int:task_id>/edit", methods=["GET", "POST"])
def edit_task(task_id):
    """任务编辑统一在导图页完成，此入口重定向过去。"""
    db.get_or_404(Task, task_id)
    return redirect(url_for("mindmap"))


@app.post("/tasks/<int:task_id>/status")
def update_task_status(task_id):
    task = db.get_or_404(Task, task_id)
    payload = request.get_json(silent=True) or request.form
    task.status = payload.get("status") or "todo"
    db.session.commit()
    sync_database()
    if request.is_json:
        return jsonify({"ok": True})
    next_url = payload.get("next", "")
    if next_url and str(next_url).startswith("/"):
        return redirect(next_url)
    return redirect(request.referrer or url_for("index"))


@app.post("/tasks/<int:task_id>/delete")
def delete_task(task_id):
    task = db.get_or_404(Task, task_id)
    for attachment in task.attachments:
        storage.delete(attachment.object_key)
    db.session.delete(task)
    db.session.commit()
    sync_database()
    flash("任务已删除", "success")
    if request.is_json:
        return jsonify({"ok": True})
    payload = request.get_json(silent=True) or request.form
    next_url = payload.get("next", "")
    if next_url and str(next_url).startswith("/"):
        return redirect(next_url)
    return redirect(url_for("index"))


@app.route("/notes/new", methods=["GET", "POST"])
def new_note():
    if request.method == "POST":
        note = Note(title=request.form["title"].strip(), content=sanitize_html(request.form.get("content")), task_id=request.form.get("task_id", type=int) or None)
        db.session.add(note)
        db.session.commit()
        sync_database()
        flash("笔记已保存", "success")
        return redirect(url_for("note_detail", note_id=note.id))
    return render_template("note_form.html", note=None, tasks=Task.query.order_by(Task.title).all(), selected_task_id=request.args.get("task_id", type=int))


@app.route("/notes/<int:note_id>")
def note_detail(note_id):
    note = db.get_or_404(Note, note_id)
    return render_template("note_detail.html", note=note, available_files=Attachment.query.order_by(Attachment.original_name).all(), tasks=Task.query.order_by(Task.title).all())


@app.route("/notes/<int:note_id>/edit", methods=["GET", "POST"])
def edit_note(note_id):
    note = db.get_or_404(Note, note_id)
    if request.method == "POST":
        note.title = request.form["title"].strip()
        note.content = sanitize_html(request.form.get("content"))
        note.task_id = request.form.get("task_id", type=int) or None
        db.session.commit()
        sync_database()
        flash("笔记已更新", "success")
        return redirect(url_for("note_detail", note_id=note.id))
    return render_template("note_form.html", note=note, tasks=Task.query.order_by(Task.title).all(), selected_task_id=note.task_id)


@app.post("/notes/<int:note_id>/delete")
def delete_note(note_id):
    note = db.session.get(Note, note_id)
    if note:  # 已不存在时静默成功（幂等），避免重复提交报 404
        for attachment in note.attachments:
            storage.delete(attachment.object_key)
        db.session.delete(note)
        db.session.commit()
        sync_database()
    flash("笔记已删除", "success")
    # 来源页若是被删笔记自己的详情页，则跳回工作台，避免跳转到已失效的地址
    referrer = request.referrer or ""
    if f"/notes/{note_id}" in referrer:
        return redirect(url_for("index"))
    return redirect(referrer or url_for("index"))


@app.route("/notes/<int:note_id>/polish", methods=["POST"])
def polish_note(note_id):
    """AI 润色已有笔记：润色后直接保存到数据库。"""
    note = db.get_or_404(Note, note_id)
    payload = request.get_json(silent=True) or {}
    raw = bleach.clean(payload.get("content") or note.content, tags=[], strip=True).strip()
    if not raw:
        return jsonify({"error": "笔记内容为空，无法润色"}), 400
    prompt = (
        "请润色以下笔记内容：修正错别字和标点，理顺语句，让表达更清晰专业，"
        "保留原有结构和要点，不要遗漏信息，不要添加知识库之外的新结论。直接输出润色后的全文，不要解释。\n\n" + raw
    )
    try:
        polished = ai.chat([{"role": "user", "content": prompt}], temperature=0.3)
    except Exception as error:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(error)}), 502
    paragraphs = [line.strip() for line in polished.splitlines() if line.strip()]
    html = "".join(f"<p>{line}</p>" for line in paragraphs)
    if not html.strip():
        # AI 返回空结果时绝不能覆盖原笔记，保底返回错误
        return jsonify({"error": "AI 返回了空内容，为保护原笔记未做任何修改"}), 502
    note.content = sanitize_html(html)
    note.updated_at = datetime.utcnow()
    db.session.commit()
    sync_database()
    return jsonify({"html": note.content, "text": polished})


@app.post("/polish")
def polish_text():
    """通用 AI 润色：编辑器中的草稿内容（尚未保存为笔记）润色后返回，不落库。"""
    payload = request.get_json(silent=True) or {}
    raw = bleach.clean(payload.get("content") or "", tags=[], strip=True).strip()
    if not raw:
        return jsonify({"error": "内容为空，无法润色"}), 400
    prompt = (
        "请润色以下笔记内容：修正错别字和标点，理顺语句，让表达更清晰专业，"
        "保留原有结构和要点，不要遗漏信息。直接输出润色后的全文，不要解释。\n\n" + raw
    )
    try:
        polished = ai.chat([{"role": "user", "content": prompt}], temperature=0.3)
    except Exception as error:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(error)}), 502
    paragraphs = [line.strip() for line in polished.splitlines() if line.strip()]
    html = "".join(f"<p>{line}</p>" for line in paragraphs)
    if not html.strip():
        return jsonify({"error": "AI 返回了空内容，请重试"}), 502
    return jsonify({"html": html, "text": polished})


# ---- 知识库 ----
@app.get("/knowledge")
def knowledge():
    docs = KnowledgeDoc.query.order_by(KnowledgeDoc.created_at.desc()).all()
    return render_template("knowledge.html", docs=docs, ai_ready=ai.enabled())


@app.post("/knowledge/upload")
def knowledge_upload():
    try:
        return _knowledge_upload_inner()
    except Exception as error:  # 兜底：任何异常都转为页面提示，避免 500
        import traceback
        traceback.print_exc()
        db.session.rollback()
        flash(f"导入失败：{error}", "error")
        return redirect(url_for("knowledge"))


def _knowledge_upload_inner():
    title = (request.form.get("title") or "").strip()
    pasted = (request.form.get("content") or "").strip()
    content = ""
    filename = ""
    upload = request.files.get("file")
    if upload and upload.filename:
        filename = upload.filename
        if not title:
            title = os.path.splitext(filename)[0]
        lower = filename.lower()
        if lower.endswith(".docx"):
            import docx
            document = docx.Document(BytesIO(upload.read()))
            content = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        elif lower.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(upload.read()))
            pages = []
            for page in reader.pages:
                text = page.extract_text() or ""
                if text.strip():
                    pages.append(text)
            content = "\n".join(pages)
            if not content.strip():
                flash("该 PDF 提取不到文字（可能是扫描件/图片版），请改用可复制文字的 PDF 或粘贴文本", "error")
                return redirect(url_for("knowledge"))
        elif lower.endswith((".txt", ".md")):
            content = upload.read().decode("utf-8", errors="ignore")
        else:
            flash("仅支持 .txt / .md / .docx / .pdf 文件，或直接粘贴文本", "error")
            return redirect(url_for("knowledge"))
    elif pasted:
        content = pasted
        if not title:
            title = "未命名文档"
    if not content.strip():
        flash("请上传文件或粘贴文本内容", "error")
        return redirect(url_for("knowledge"))
    doc = KnowledgeDoc(title=title[:200], filename=filename, content=content)
    db.session.add(doc)
    db.session.commit()
    try:
        embed_document(doc)
    except Exception as error:
        db.session.delete(doc)
        db.session.commit()
        raise error
    sync_database()
    flash(f"已导入「{doc.title}」（{len(doc.chunks)} 个知识块）", "success")
    return redirect(url_for("knowledge"))


@app.post("/knowledge/doc/<int:doc_id>/delete")
def knowledge_delete(doc_id):
    doc = db.get_or_404(KnowledgeDoc, doc_id)
    db.session.delete(doc)
    db.session.commit()
    sync_database()
    flash("文档已删除", "success")
    return redirect(url_for("knowledge"))


@app.post("/knowledge/ask")
def knowledge_ask():
    question = (request.get_json(silent=True) or {}).get("question", "").strip()
    if not question:
        return jsonify({"error": "请输入问题"}), 400
    try:
        result = ask_knowledge(question)
    except Exception as error:
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(error)}), 502
    return jsonify(result)


@app.route("/files")
def file_manager():
    folder_id = request.args.get("folder_id", type=int)
    folder = db.session.get(Folder, folder_id) if folder_id else None
    if folder_id and not folder:
        return redirect(url_for("file_manager"))
    folders = Folder.query.filter_by(parent_id=folder_id).order_by(Folder.name).all()
    if folder:
        files = Attachment.query.filter_by(folder_id=folder.id).order_by(Attachment.created_at.desc()).all()
    else:
        files = Attachment.query.order_by(Attachment.created_at.desc()).all()
    all_folders = Folder.query.order_by(Folder.name).all()
    children_map = {}
    for entry in all_folders:
        children_map.setdefault(entry.parent_id, []).append(entry)

    def folder_path(entry):
        parts = []
        while entry:
            parts.append(entry.name)
            entry = entry.parent
        return " / ".join(reversed(parts))

    folder_paths = {entry.id: folder_path(entry) for entry in all_folders}
    breadcrumbs = []
    current = folder
    while current:
        breadcrumbs.append(current)
        current = current.parent
    breadcrumbs = list(reversed(breadcrumbs))
    expanded_ids = {entry.id for entry in breadcrumbs}
    return render_template("file_manager.html", folder=folder, folders=folders, files=files, all_folders=all_folders,
                           folder_paths=folder_paths, children_map=children_map, breadcrumbs=breadcrumbs,
                           expanded_ids=expanded_ids)


@app.post("/files/folders")
def create_folder():
    name = request.form.get("name", "").strip()
    parent_id = request.form.get("parent_id", type=int) or None
    if not name:
        flash("请输入文件夹名称", "error")
    else:
        db.session.add(Folder(name=name, parent_id=parent_id))
        db.session.commit()
        sync_database()
        flash("文件夹已创建", "success")
    return redirect(url_for("file_manager", folder_id=parent_id))


def folder_has_content(folder):
    if Attachment.query.filter_by(folder_id=folder.id).first() is not None:
        return True
    children = Folder.query.filter_by(parent_id=folder.id).all()
    return bool(children) or any(folder_has_content(child) for child in children)


@app.post("/files/folders/<int:folder_id>/rename")
def rename_folder(folder_id):
    folder = db.get_or_404(Folder, folder_id)
    name = request.form.get("name", "").strip()
    if name:
        folder.name = name
        db.session.commit()
        sync_database()
        flash("文件夹已改名", "success")
    else:
        flash("文件夹名称不能为空", "error")
    return redirect(url_for("file_manager", folder_id=folder.parent_id))


@app.post("/files/folders/<int:folder_id>/delete")
def delete_folder(folder_id):
    folder = db.get_or_404(Folder, folder_id)
    parent_id = folder.parent_id
    if folder_has_content(folder):
        flash("文件夹或其子文件夹不为空，不能删除", "error")
    else:
        db.session.delete(folder)
        db.session.commit()
        sync_database()
        flash("文件夹已删除", "success")
    return redirect(url_for("file_manager", folder_id=parent_id))


@app.post("/files/<int:attachment_id>/move")
def move_file(attachment_id):
    attachment = db.get_or_404(Attachment, attachment_id)
    payload = request.get_json(silent=True) or request.form
    raw_folder_id = payload.get("folder_id") if hasattr(payload, "get") else None
    folder_id = int(raw_folder_id) if raw_folder_id not in (None, "", 0, "0") else None
    if folder_id and not db.session.get(Folder, folder_id):
        return jsonify({"error": "文件夹不存在"}), 404
    attachment.folder_id = folder_id or None
    db.session.commit()
    sync_database()
    if not request.is_json:
        flash("文件已移动", "success")
        return redirect(request.referrer or url_for("file_manager"))
    return jsonify({"ok": True})


@app.post("/files/<int:attachment_id>/delete")
def delete_file(attachment_id):
    attachment = db.get_or_404(Attachment, attachment_id)
    storage.delete(attachment.object_key)
    db.session.delete(attachment)
    db.session.commit()
    sync_database()
    flash("文件已删除", "success")
    return redirect(request.referrer or url_for("file_manager"))


@app.get("/attachments/check-name")
def check_attachment_name():
    filename = request.args.get("filename", "").strip()
    duplicate = bool(filename and Attachment.query.filter(func.lower(Attachment.original_name) == filename.lower()).first())
    return jsonify({"duplicate": duplicate})


@app.post("/attachments/link")
def link_attachment():
    attachment_id = request.form.get("attachment_id", type=int)
    if not attachment_id:
        flash("请选择文件", "error")
        return redirect(request.referrer or url_for("file_manager"))
    attachment = db.get_or_404(Attachment, attachment_id)
    task_id = request.form.get("task_id", type=int)
    note_id = request.form.get("note_id", type=int)
    if not task_id and not note_id:
        flash("请选择任务或笔记", "error")
    else:
        if task_id:
            db.get_or_404(Task, task_id)
            attachment.task_id = task_id
        if note_id:
            db.get_or_404(Note, note_id)
            attachment.note_id = note_id
        db.session.commit()
        sync_database()
        flash("文件已关联", "success")
    return redirect(request.referrer or url_for("file_manager"))


@app.post("/attachments/<int:attachment_id>/unlink")
def unlink_attachment(attachment_id):
    attachment = db.get_or_404(Attachment, attachment_id)
    target = request.form.get("target")
    if target == "task" and attachment.task_id:
        attachment.task_id = None
        db.session.commit()
        sync_database()
        flash("已解除与任务的关联，文件仍保留在文件管理中", "success")
    elif target == "note" and attachment.note_id:
        attachment.note_id = None
        db.session.commit()
        sync_database()
        flash("已解除与笔记的关联，文件仍保留在文件管理中", "success")
    else:
        flash("没有可解除的关联", "error")
    return redirect(request.referrer or url_for("file_manager"))


def mindmap_tree_data():
    def task_node(task):
        return {
            "id": task.id,
            "title": task.title,
            "description": task.description or "",
            "status": task.status,
            "priority": task.priority,
            "due_date": task.due_date.strftime("%Y-%m-%d") if task.due_date else "",
            "notes": [{"id": note.id, "title": note.title} for note in task.notes],
            "attachments": [{"id": att.id, "name": att.original_name} for att in task.attachments],
            "children": [task_node(child) for child in task.children],
        }

    roots = Task.query.filter_by(parent_id=None).order_by(Task.created_at).all()
    return [task_node(task) for task in roots]


@app.context_processor
def inject_static_version():
    def static_version(filename):
        path = os.path.join(app.static_folder or "static", filename)
        try:
            return str(int(os.path.getmtime(path)))
        except OSError:
            return "0"
    return dict(static_version=static_version)


@app.get("/mindmap")
def mindmap():
    tree = mindmap_tree_data()
    notes = Note.query.order_by(Note.updated_at.desc()).all()
    attachments = Attachment.query.order_by(Attachment.created_at.desc()).all()
    return render_template("mindmap.html", tree=tree, notes=notes, attachments=attachments)


@app.get("/mindmap/data")
def mindmap_data():
    return jsonify({"tree": mindmap_tree_data()})


@app.post("/mindmap/tasks/create")
def mindmap_create_task():
    payload = request.get_json(silent=True) or request.form
    title = (payload.get("title") or "").strip()
    raw_parent = payload.get("parent_id")
    parent_id = int(raw_parent) if raw_parent not in (None, "", 0, "0") else None
    if not title:
        if request.is_json:
            return jsonify({"error": "请填写节点名称"}), 400
        flash("请填写节点名称", "error")
    else:
        if parent_id:
            db.get_or_404(Task, parent_id)
        task = Task(title=title, parent_id=parent_id)
        db.session.add(task)
        db.session.commit()
        sync_database()
        if request.is_json:
            return jsonify({"ok": True, "id": task.id})
        flash("节点已创建", "success")
    return redirect(url_for("mindmap"))


@app.post("/mindmap/tasks/<int:task_id>/move")
def mindmap_move_task(task_id):
    task = db.get_or_404(Task, task_id)
    payload = request.get_json(silent=True) or request.form
    raw_parent = payload.get("parent_id")
    new_parent_id = int(raw_parent) if raw_parent not in (None, "", 0, "0") else None
    if new_parent_id == task.id:
        return jsonify({"error": "不能移动到自身"}), 400
    if new_parent_id:
        ancestor = db.session.get(Task, new_parent_id)
        while ancestor:
            if ancestor.id == task.id:
                return jsonify({"error": "不能移动到自己的子节点下"}), 400
            ancestor = ancestor.parent
        task.parent_id = new_parent_id
    else:
        task.parent_id = None
    db.session.commit()
    sync_database()
    return jsonify({"ok": True})


@app.post("/mindmap/tasks/<int:task_id>/delete-node")
def mindmap_delete_node(task_id):
    """仅删除当前节点，子节点提升一级（模仿 MindNow 删除逻辑）"""
    with db.session.no_autoflush:
        task = db.get_or_404(Task, task_id)
        old_parent_id = task.parent_id
        children_ids = [child.id for child in Task.query.filter_by(parent_id=task_id).all()]
    db.session.expunge_all()
    db.session.execute(db.update(Task).where(Task.parent_id == task_id).values(parent_id=old_parent_id))
    db.session.commit()
    task = db.session.get(Task, task_id)
    db.session.delete(task)
    db.session.commit()
    sync_database()
    return jsonify({"ok": True})


@app.post("/mindmap/tasks/<int:task_id>/rename")
def mindmap_rename_task(task_id):
    task = db.get_or_404(Task, task_id)
    payload = request.get_json(silent=True) or request.form
    title = (payload.get("title") or "").strip()
    if title:
        task.title = title
        db.session.commit()
        sync_database()
        if request.is_json:
            return jsonify({"ok": True})
        flash("节点已重命名", "success")
    return redirect(url_for("mindmap"))


@app.post("/mindmap/tasks/<int:task_id>/update")
def mindmap_update_task(task_id):
    task = db.get_or_404(Task, task_id)
    payload = request.get_json(silent=True) or request.form
    title = (payload.get("title") or "").strip()
    if title:
        task.title = title
    task.description = sanitize_html(payload.get("description"))
    task.status = payload.get("status") or "todo"
    task.priority = payload.get("priority") or "medium"
    task.due_date = parse_date(payload.get("due_date"))
    db.session.commit()
    sync_database()
    if request.is_json:
        return jsonify({"ok": True})
    flash("任务已更新", "success")
    return redirect(url_for("mindmap"))


@app.post("/mindmap/tasks/<int:task_id>/link")
def mindmap_link(task_id):
    task = db.get_or_404(Task, task_id)
    payload = request.get_json(silent=True) or request.form
    kind = payload.get("kind")
    try:
        item_id = int(payload.get("item_id") or 0)
    except (TypeError, ValueError):
        item_id = 0
    if kind == "note" and item_id:
        note = db.get_or_404(Note, item_id)
        note.task_id = task.id
    elif kind == "attachment" and item_id:
        attachment = db.get_or_404(Attachment, item_id)
        attachment.task_id = task.id
    else:
        if request.is_json:
            return jsonify({"error": "请选择要关联的内容"}), 400
        flash("请选择要关联的内容", "error")
        return redirect(url_for("mindmap"))
    db.session.commit()
    sync_database()
    if request.is_json:
        return jsonify({"ok": True})
    flash("已关联", "success")
    return redirect(url_for("mindmap"))


@app.post("/mindmap/tasks/<int:task_id>/unlink")
def mindmap_unlink(task_id):
    task = db.get_or_404(Task, task_id)
    payload = request.get_json(silent=True) or request.form
    kind = payload.get("kind")
    try:
        item_id = int(payload.get("item_id") or 0)
    except (TypeError, ValueError):
        item_id = 0
    if kind == "note" and item_id:
        note = db.get_or_404(Note, item_id)
        if note.task_id == task.id:
            note.task_id = None
    elif kind == "attachment" and item_id:
        attachment = db.get_or_404(Attachment, item_id)
        if attachment.task_id == task.id:
            attachment.task_id = None
    db.session.commit()
    sync_database()
    if request.is_json:
        return jsonify({"ok": True})
    flash("已解除关联", "success")
    return redirect(url_for("mindmap"))


@app.post("/attachments/upload")
def upload_attachment():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        flash("请选择附件", "error")
        return redirect(request.referrer or url_for("index"))
    if not storage.enabled:
        flash("附件上传需要先配置阿里云 OSS 环境变量", "error")
        return redirect(request.referrer or url_for("index"))
    task_id = request.form.get("task_id", type=int)
    note_id = request.form.get("note_id", type=int)
    folder_id = request.form.get("folder_id", type=int)
    allow_duplicate = request.form.get("allow_duplicate") == "1"
    duplicate = Attachment.query.filter(func.lower(Attachment.original_name) == uploaded.filename.strip().lower()).first()
    if duplicate and not allow_duplicate:
        flash(f"已存在同名文件“{uploaded.filename}”，如需继续上传请确认重复上传。", "error")
        return redirect(request.referrer or url_for("file_manager"))
    safe_name = re.sub(r"[^\w.\- ]", "_", uploaded.filename)[:180]
    key = f"attachments/{uuid.uuid4().hex}-{safe_name}"
    try:
        uploaded.stream.seek(0)
        uploaded_size = storage.upload(uploaded, key)
        attachment = Attachment(original_name=uploaded.filename, object_key=key, content_type=uploaded.content_type or "application/octet-stream", size=uploaded_size, task_id=task_id, note_id=note_id, folder_id=folder_id)
        db.session.add(attachment)
        db.session.commit()
        sync_database()
        flash("附件已上传", "success")
    except (BotoCoreError, ClientError, RuntimeError) as error:
        db.session.rollback()
        flash(f"附件上传失败：{error}", "error")
    return redirect(request.referrer or url_for("index"))


@app.get("/attachments/<int:attachment_id>")
def download_attachment(attachment_id):
    attachment = db.get_or_404(Attachment, attachment_id)
    try:
        data = storage.download(attachment.object_key)
    except (BotoCoreError, ClientError) as error:
        flash(f"附件读取失败：{error}", "error")
        return redirect(request.referrer or url_for("index"))
    return send_file(BytesIO(data), as_attachment=True, download_name=attachment.original_name, mimetype=attachment.content_type)


@app.get("/health")
def health():
    return {"status": "ok"}


with app.app_context():
    restore_database()
    db.create_all()
    migrate_schema()


if __name__ == "__main__":
    app.run(debug=os.getenv("FLASK_DEBUG", "0") == "1", host="0.0.0.0", port=int(os.getenv("PORT", "5000")))
